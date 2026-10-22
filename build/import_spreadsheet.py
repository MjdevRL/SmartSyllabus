"""One-time import of 'Final University List.xlsx' into content/universities/*.json.

Rules: every value is copied exactly as written (only stray spaces at the start/end of
cells are removed). Nothing is corrected. Questionable details become review flags.
Usage: python3 -I build/import_spreadsheet.py "<path to xlsx>" content
"""
import json, re, sys, unicodedata, pathlib
import openpyxl

src, out = sys.argv[1], pathlib.Path(sys.argv[2])
ws = openpyxl.load_workbook(src).active

def slug(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower().replace('&', ' and ')
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')[:80].strip('-')

def cell(r, c):
    v = ws.cell(r, c).value
    return v.strip() if isinstance(v, str) else v

GENERAL = [
    "Costs: the currency isn't stated, and it isn't recorded whether the figure is tuition only or tuition plus living costs.",
    "The website column links to the university home page. Add the official course page.",
    "Figures come from personal research and haven't been checked against official sources yet.",
]
SPECIFIC = {
    'Southampton': ["Location says 'Oxford Road, Manchester', which looks like the University of Manchester's address. Southampton's main campus is in Southampton.",
                    "Check whether this course interviews applicants."],
    'Nortingham': ["Name is spelled 'Nortingham'. Did you mean Nottingham?",
                   "Location (South Kensington, London) and notes match the Imperial College London row, and the ranking reason says 'Shares Imperial's outstanding reputation'. This row may be a copy of the Imperial entry that was renamed.",
                   "Course says 'without honours'. Check the exact programme title on the official course page."],
    'UCL': ["Check whether this course requires the ESAT and an interview."],
    'TUM Asia': ["Country says Germany but the location is in Singapore (SIT Punggol campus). Record where the degree is awarded and where you'd study separately.",
                 "Website is SIT's home page, not TUM Asia's."],
    'RWU': ["Rank is written as '16rth' (read here as 16th)."],
}
US = ["US universities generally don't publish IB score cut-offs, so treat the 41-45 range as an estimate.",
      "Testing policies (SAT/ACT required, optional or 'test-flexible') change between application cycles. Check the official admissions page.",
      "US applications usually involve supplemental essays and recommendations as well as a main essay."]
DE = ["German admission usually depends on your IB subject combination and grades. Check whether the SAT is required at all, and treat the IB score range as an estimate."]
SG = ["Check whether IELTS is required for IB applicants. It may depend on your IB English results."]

(out / 'universities').mkdir(parents=True, exist_ok=True)
count = 0
for r in range(2, ws.max_row + 1):
    name = cell(r, 1)
    if not name or cell(r, 2) is None:
        continue
    if str(name).startswith('Extra Exams'):
        break
    rank_text = str(cell(r, 12) or '')
    m = re.match(r'\s*(\d+)', rank_text)
    country = cell(r, 7)
    flags = list(SPECIFIC.get(name, []))
    if country == 'USA': flags += US
    if country == 'Germany' and name != 'TUM Asia': flags += DE
    if country == 'Singapore': flags += SG
    flags += GENERAL
    entry = {
        'university': name,
        'course': cell(r, 2),
        'country': country,
        'location': cell(r, 8),
        'website': cell(r, 14),
        'course_page': '',
        'ib_score_range': str(cell(r, 3) or ''),
        'extra_exams': cell(r, 4),
        'recognition_score': cell(r, 5),
        'roi_score': cell(r, 6),
        'yearly_cost': cell(r, 9),
        'cost_currency': '',
        'cost_includes': '',
        'first_salary': cell(r, 10),
        'other_notes': cell(r, 11),
        'personal_rank': rank_text,
        'personal_rank_number': int(m.group(1)) if m else None,
        'rank_justification': cell(r, 13),
        'fact_status': 'unverified',
        'source_url': '',
        'last_verified': '',
        'review_flags': flags,
        'status': 'published',
        'notices': [],
        'spreadsheet_row': r,
    }
    fn = out / 'universities' / f"{slug(name + ' ' + entry['course'])}.json"
    fn.write_text(json.dumps(entry, indent=2, ensure_ascii=False) + '\n')
    count += 1

# The small "extra exams" table further down the sheet: your personal score targets.
targets = []
start = next(r for r in range(1, ws.max_row + 1) if str(cell(r, 1) or '').startswith('Extra Exams'))
for r in range(start + 1, ws.max_row + 1):
    if cell(r, 1) is None:
        continue
    targets.append({'exam': cell(r, 1), 'upper_target': cell(r, 2), 'lower_target': cell(r, 3)})
for t in targets:
    if str(t['exam']).startswith('IELTS') and t['upper_target'] == 8.7:
        t['review_flag'] = 'IELTS scores go up in steps of 0.5 to a maximum of 9, so 8.7 isn\'t a possible score.'
    if str(t['exam']).startswith('ESAT'):
        t['review_flag'] = "Check these targets against the ESAT's official score scale."
(out / 'settings').mkdir(parents=True, exist_ok=True)
(out / 'settings' / 'test-targets.json').write_text(json.dumps({'note': 'Imported from rows %d-%d of your spreadsheet. Personal targets, not entry requirements.' % (start, ws.max_row), 'targets': targets}, indent=2, ensure_ascii=False) + '\n')
print(count, 'universities;', len(targets), 'test targets')
